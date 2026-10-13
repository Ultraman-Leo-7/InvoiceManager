from pathlib import Path
from types import SimpleNamespace

import fitz
import pytest
from openpyxl import load_workbook

import app
import invoice_extract
from purchase_tracker import add_purchase, set_manual_match


@pytest.fixture
def invoice_folder(tmp_path, monkeypatch):
    monkeypatch.setattr(app, "BASE_DIR", tmp_path)
    monkeypatch.setattr(app, "DB_PATH", tmp_path / "test.db")
    monkeypatch.setattr(app, "EXPORT_PATH", tmp_path / "test.xlsx")
    monkeypatch.setattr(app, "LOCAL_BACKUP_DIR", tmp_path / "backups")

    def parse_test_invoice(path):
        return {
            "文件名": path.name,
            "发票号码": app.sha256_file(path)[:20],
            "项目名称": "测试商品",
            "价税合计": 12.5,
            "提取状态": "正常",
        }

    def unexpected_dialog(*args, **kwargs):
        pytest.fail("Folder scanning must leave confirmation to the GUI")

    monkeypatch.setattr(app, "parse_invoice", parse_test_invoice)
    monkeypatch.setattr(app.messagebox, "askyesno", unexpected_dialog)
    app.init_db()
    folder = tmp_path / "发票"
    folder.mkdir(exist_ok=True)
    return folder


def make_pdf(folder, name, text="Test invoice"):
    path = folder / name
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), text)
        document.save(path)
    return path


def set_arrival_times(monkeypatch, folder, arrival_times):
    original_stat = Path.stat

    def file_stat(path, *args, **kwargs):
        info = original_stat(path, *args, **kwargs)
        if path.parent == folder and path.name in arrival_times:
            fields = {name: getattr(info, name) for name in dir(info) if name.startswith("st_")}
            fields["st_birthtime"] = arrival_times[path.name]
            return SimpleNamespace(**fields)
        return info

    monkeypatch.setattr(Path, "stat", file_stat)


def test_first_scan_only_detects_duplicates_and_prefers_earliest_copy(invoice_folder, monkeypatch):
    original = make_pdf(invoice_folder, "z-old.pdf")
    content = original.read_bytes()
    (invoice_folder / "a-new.pdf").write_bytes(content)
    (invoice_folder / "m-middle.pdf").write_bytes(content)
    set_arrival_times(monkeypatch, invoice_folder, {"z-old.pdf": 10, "m-middle.pdf": 20, "a-new.pdf": 30})

    stats = app.sync_folder()

    assert stats["failed"] == 0
    assert stats["unique"] == 1
    assert stats["total"] == 3
    assert stats["duplicates"] == [
        {"digest": app.sha256_file(original), "keep": "z-old.pdf", "remove": ["m-middle.pdf", "a-new.pdf"]}
    ]
    assert sorted(path.name for path in invoice_folder.glob("*.pdf")) == ["a-new.pdf", "m-middle.pdf", "z-old.pdf"]
    assert all(path.read_bytes() == content for path in invoice_folder.glob("*.pdf"))
    rows = app.all_active_rows()
    assert len(rows) == 1
    assert rows[0]["filename"] == "z-old.pdf"
    assert app.sync_folder()["duplicates"] == stats["duplicates"]
    assert len(list(invoice_folder.glob("*.pdf"))) == 3


@pytest.mark.parametrize("copy_name", ["a-new.pdf", "z-new.pdf"])
def test_new_copy_preserves_existing_database_filename_until_decision(invoice_folder, copy_name):
    original = make_pdf(invoice_folder, "m-original.pdf")
    app.sync_folder()
    digest = app.all_active_rows()[0]["digest"]
    (invoice_folder / copy_name).write_bytes(original.read_bytes())

    stats = app.sync_folder()

    assert original.exists()
    assert (invoice_folder / copy_name).exists()
    assert stats["duplicates"] == [{"digest": digest, "keep": "m-original.pdf", "remove": [copy_name]}]
    rows = app.all_active_rows()
    assert len(rows) == 1
    assert rows[0]["digest"] == digest
    assert rows[0]["filename"] == "m-original.pdf"


def test_changed_content_stays_separate_and_is_never_offered_for_deletion(invoice_folder):
    original = make_pdf(invoice_folder, "a.pdf", "First invoice")
    app.sync_folder()
    changed = make_pdf(invoice_folder, "b.pdf", "Different invoice")

    stats = app.sync_folder()

    assert stats["duplicates"] == []
    assert stats["unique"] == 2
    assert original.exists()
    assert changed.exists()
    rows = app.all_active_rows()
    assert [row["filename"] for row in rows] == ["a.pdf", "b.pdf"]
    assert len({row["digest"] for row in rows}) == 2


@pytest.mark.parametrize("content", [b"", b"%PDF-1.7\nunfinished", b"not a PDF"])
def test_empty_or_malformed_pdfs_are_not_offered_for_cleanup(invoice_folder, monkeypatch, content):
    monkeypatch.setattr(app, "parse_invoice", invoice_extract.parse_invoice)
    (invoice_folder / "a.pdf").write_bytes(content)
    (invoice_folder / "b.pdf").write_bytes(content)

    stats = app.sync_folder()

    assert stats["duplicates"] == []
    assert stats["failed"] == 0
    assert stats["abnormal_count"] == 2
    assert (app.abnormal_dir() / "a.pdf").read_bytes() == content
    assert (app.abnormal_dir() / "b.pdf").read_bytes() == content
    assert not list(invoice_folder.glob("*.pdf"))


@pytest.mark.parametrize("delete", [True, False], ids=["confirm-delete", "decline-and-prefix"])
def test_duplicate_decision_preserves_manual_data_and_single_excel_row(invoice_folder, monkeypatch, delete):
    original = make_pdf(invoice_folder, "original.pdf")
    app.sync_folder()
    digest = app.all_active_rows()[0]["digest"]
    app.set_confirmed_many([digest], True)
    app.set_note(digest, "人工复核完成")
    with app.connect_db() as conn:
        purchase_id = add_purchase(conn, "人工指定订单", "99.00", False, 0)
        set_manual_match(conn, purchase_id, "商品", digest)

    renamed = invoice_folder / "z-renamed.pdf"
    original.rename(renamed)
    content = renamed.read_bytes()
    (invoice_folder / "a-copy.pdf").write_bytes(content)
    (invoice_folder / "b-copy.pdf").write_bytes(content)
    set_arrival_times(monkeypatch, invoice_folder, {"z-renamed.pdf": 10, "a-copy.pdf": 20, "b-copy.pdf": 30})
    stats = app.sync_folder()
    assert len(list(invoice_folder.glob("*.pdf"))) == 3

    result = app.handle_duplicate_files(stats["duplicates"], delete=delete)

    assert result["errors"] == []
    assert result["stats"]["duplicates"] == []
    assert result["stats"]["unique"] == 1
    if delete:
        assert sorted(result["deleted"]) == ["a-copy.pdf", "b-copy.pdf"]
        assert result["renamed"] == []
        assert sorted(path.name for path in invoice_folder.glob("*.pdf")) == ["z-renamed.pdf"]
    else:
        assert result["deleted"] == []
        assert sorted(result["renamed"]) == [("a-copy.pdf", "哈希副本-a-copy.pdf"), ("b-copy.pdf", "哈希副本-b-copy.pdf")]
        assert sorted(path.name for path in invoice_folder.glob("*.pdf")) == ["z-renamed.pdf", "哈希副本-a-copy.pdf", "哈希副本-b-copy.pdf"]
    assert all(path.read_bytes() == content for path in invoice_folder.glob("*.pdf"))
    rows = app.all_active_rows()
    assert len(rows) == 1
    assert rows[0]["digest"] == digest
    assert rows[0]["filename"] == "z-renamed.pdf"
    assert rows[0]["confirmed"] == 1
    assert rows[0]["note"] == "人工复核完成"
    matches = app.current_match_data()
    component = matches["purchase_results"][0]["components"][0]
    assert component["manual"] is True
    assert component["matched"] is True
    assert component["invoice"]["digest"] == digest
    assert component["invoice"]["filename"] == "z-renamed.pdf"
    with app.connect_db() as conn:
        override = conn.execute("SELECT purchase_id, invoice_digest FROM purchase_invoice_overrides").fetchone()
        assert tuple(override) == (purchase_id, digest)
        assert conn.execute("SELECT COUNT(*) FROM purchases").fetchone()[0] == 1

    assert app.export_excel(["文件名", "价税合计"]) == 1
    workbook = load_workbook(app.EXPORT_PATH, data_only=True)
    try:
        sheet = workbook["发票汇总"]
        assert sheet.max_row == 2
        assert list(sheet.values)[1] == ("z-renamed.pdf", 12.5, "人工指定订单（商品，手动）", "是", "人工复核完成")
        assert workbook["购买记录"]["E2"].value == "z-renamed.pdf"
    finally:
        workbook.close()


def test_declined_copies_are_not_prompted_again_but_new_copy_is(invoice_folder):
    original = make_pdf(invoice_folder, "original.pdf")
    app.sync_folder()
    (invoice_folder / "copy.pdf").write_bytes(original.read_bytes())
    stats = app.sync_folder()

    result = app.handle_duplicate_files(stats["duplicates"], delete=False)

    assert result["errors"] == []
    assert (invoice_folder / "哈希副本-copy.pdf").exists()
    assert not (invoice_folder / "copy.pdf").exists()
    app.init_db()
    assert app.sync_folder()["duplicates"] == []
    assert not (invoice_folder / "哈希副本-哈希副本-copy.pdf").exists()

    (invoice_folder / "new-copy.pdf").write_bytes(original.read_bytes())
    new_stats = app.sync_folder()
    assert len(new_stats["duplicates"]) == 1
    assert new_stats["duplicates"][0]["keep"] == "original.pdf"
    assert new_stats["duplicates"][0]["remove"] == ["new-copy.pdf"]
    assert len(list(invoice_folder.glob("*.pdf"))) == 3


def test_decline_prefix_collision_never_overwrites_existing_file(invoice_folder):
    original = make_pdf(invoice_folder, "original.pdf")
    app.sync_folder()
    content = original.read_bytes()
    (invoice_folder / "copy.pdf").write_bytes(content)
    existing = make_pdf(invoice_folder, "哈希副本-copy.pdf", "Unrelated invoice")
    existing_content = existing.read_bytes()
    stats = app.sync_folder()

    result = app.handle_duplicate_files(stats["duplicates"], delete=False)

    assert result["errors"] == []
    assert result["deleted"] == []
    assert len(result["renamed"]) == 1
    old_name, new_name = result["renamed"][0]
    assert old_name == "copy.pdf"
    assert new_name.startswith("哈希副本-")
    assert new_name != existing.name
    assert (invoice_folder / new_name).read_bytes() == content
    assert existing.read_bytes() == existing_content
    assert original.read_bytes() == content
    assert result["stats"]["unique"] == 2
    assert result["stats"]["duplicates"] == []


def test_normal_delete_after_confirmed_cleanup_does_not_reappear_on_refresh(invoice_folder, monkeypatch):
    original = make_pdf(invoice_folder, "original.pdf")
    app.sync_folder()
    digest = app.all_active_rows()[0]["digest"]
    (invoice_folder / "copy-a.pdf").write_bytes(original.read_bytes())
    (invoice_folder / "copy-b.pdf").write_bytes(original.read_bytes())
    stats = app.sync_folder()
    app.handle_duplicate_files(stats["duplicates"], delete=True)
    assert len(list(invoice_folder.glob("*.pdf"))) == 1
    sync_results = []

    def run_sync(**kwargs):
        sync_results.append(app.sync_folder())

    view = SimpleNamespace(selected_digests=lambda: [digest], run_background_sync=run_sync)
    monkeypatch.setattr(app.messagebox, "askyesno", lambda *args, **kwargs: True)
    monkeypatch.setattr(app.messagebox, "showwarning", lambda *args, **kwargs: pytest.fail("Deletion failed"))
    app.InvoiceApp.delete_selected_invoices(view)

    assert not list(invoice_folder.glob("*.pdf"))
    assert sync_results[0]["removed"] == 1
    assert sync_results[0]["unique"] == 0
    assert not app.all_active_rows()
    assert app.sync_folder()["unique"] == 0
    with app.connect_db() as conn:
        assert conn.execute("SELECT active FROM invoices WHERE digest=?", (digest,)).fetchone()[0] == 0


@pytest.mark.parametrize("delete", [True, False])
def test_keeper_disappearing_before_decision_preserves_survivor_and_reconciles_database(invoice_folder, delete):
    original = make_pdf(invoice_folder, "original.pdf")
    app.sync_folder()
    digest = app.all_active_rows()[0]["digest"]
    copy = invoice_folder / "copy.pdf"
    copy.write_bytes(original.read_bytes())
    stats = app.sync_folder()
    original.unlink()

    result = app.handle_duplicate_files(stats["duplicates"], delete=delete)

    assert result["deleted"] == []
    assert result["renamed"] == []
    assert result["errors"]
    assert copy.exists()
    assert result["stats"]["unique"] == 1
    rows = app.all_active_rows()
    assert len(rows) == 1
    assert rows[0]["digest"] == digest
    assert rows[0]["filename"] == "copy.pdf"
    assert (invoice_folder / rows[0]["filename"]).exists()


@pytest.mark.parametrize("delete", [True, False])
def test_victim_changed_before_decision_is_preserved_and_rescanned(invoice_folder, delete):
    original = make_pdf(invoice_folder, "original.pdf")
    app.sync_folder()
    copy = invoice_folder / "copy.pdf"
    copy.write_bytes(original.read_bytes())
    stats = app.sync_folder()
    make_pdf(invoice_folder, copy.name, "Different invoice after scan")
    changed_content = copy.read_bytes()

    result = app.handle_duplicate_files(stats["duplicates"], delete=delete)

    assert result["deleted"] == []
    assert result["renamed"] == []
    assert result["errors"]
    assert original.exists()
    assert copy.read_bytes() == changed_content
    assert app.sha256_file(original) != app.sha256_file(copy)
    assert result["stats"]["unique"] == 2
    assert result["stats"]["duplicates"] == []
    assert len({row["digest"] for row in app.all_active_rows()}) == 2
